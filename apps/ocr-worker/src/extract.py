"""Extraction de texte par type de fichier
"""

import pdfplumber
import pytesseract
from docx import Document as DocxDocument
from pdf2image import convert_from_path
from PIL import Image
from pptx import Presentation

from .mime_type import file_type_from_mime_type

# En dessous de ce nombre de caractères de texte natif extrait, un PDF est considéré comme
# scanné (pas de calque texte, juste des images de pages) et bascule sur l'OCR Tesseract.
NATIVE_TEXT_MIN_LENGTH = 20
PDF_OCR_DPI = 200


def _strip_null_bytes(text: str) -> str:
    """pdfplumber peut extraire un octet NUL (`\\x00`) pour certaines ligatures mal décodées : Postgres rejette purement et simplement
    tout texte contenant un NUL dans une colonne texte, ça doit donc être nettoyé ici, à la source, plutôt que de laisser exploser l'écriture en base côté
    api-service."""
    return text.replace("\x00", "")


def extract_text_from_pdf(pdf_path: str) -> str:
    with pdfplumber.open(pdf_path) as pdf:
        native_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    if len(native_text.strip()) >= NATIVE_TEXT_MIN_LENGTH:
        return native_text

    # Pas de calque texte exploitable : PDF scanné, on rend chaque page en image et on OCRise.
    pages = convert_from_path(pdf_path, dpi=PDF_OCR_DPI)
    return "\n".join(pytesseract.image_to_string(page) for page in pages)


def extract_text_from_image(image_path: str) -> str:
    with Image.open(image_path) as image:
        return pytesseract.image_to_string(image)


def extract_text_from_pptx(pptx_path: str) -> str:
    presentation = Presentation(pptx_path)
    chunks: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


def extract_text_from_docx(docx_path: str) -> str:
    document = DocxDocument(docx_path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def extract_text_from_markdown(markdown_path: str) -> str:
    with open(markdown_path, encoding="utf-8", errors="replace") as f:
        return f.read()


def extract(mime_type: str, local_path: str) -> tuple[str, bool]:
    """Retourne (texte extrait, supported). `supported=False` signifie qu'aucun extracteur ne
    prend en charge ce mimeType"""
    file_type = file_type_from_mime_type(mime_type)

    if file_type == "pdf":
        return _strip_null_bytes(extract_text_from_pdf(local_path)), True
    if file_type == "image":
        return _strip_null_bytes(extract_text_from_image(local_path)), True
    if file_type == "pptx":
        return _strip_null_bytes(extract_text_from_pptx(local_path)), True
    if file_type == "docx":
        return _strip_null_bytes(extract_text_from_docx(local_path)), True
    if file_type == "markdown":
        return _strip_null_bytes(extract_text_from_markdown(local_path)), True

    return "", False
